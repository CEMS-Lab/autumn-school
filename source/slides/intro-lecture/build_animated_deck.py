"""Extend the retained editable v7 with five original, embedded teaching clips.

Run from the course checkout with Python 3.10+, python-pptx and FFmpeg available:
    python source/slides/intro-lecture/build_animated_deck.py

The v7 source deck and the original animation files are read-only inputs.
Created by Allamaprabhu Ani, CEMS-Lab, UKACM Autumn School 2026.
"""
from __future__ import annotations

import hashlib
import json
import re
import shutil
import subprocess
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Pt

ROOT = Path(__file__).resolve().parents[3]
SOURCE = ROOT / "output/intro-lecture-20260909/phast_lecture_introduction_v7.pptx"
OUT = ROOT / "output/lecture-20260910"
BLUE, ORANGE, INK, MUTED = "245A81", "B85C20", "17191C", "525D67"
CREDIT = ("Created by Allamaprabhu Ani, CEMS-Lab, for the UKACM Autumn School 2026. "
          "Presented by Sathiskumar A. Ponnusami, Queen Mary University of London.")

CLIPS = [
    dict(key="phase_field_band", title="The length scale controls the crack band",
         movie="assets/animations/mechanics/phase_field_band.mp4",
         poster="assets/animations/mechanics/phase_field_band_poster.png",
         section="L1 · 13–26 min", pause=15.0,
         point="Larger length scale\n\nWider diffuse band\n\nFiner mesh resolves smaller bands",
         question="How should the mesh change when the regularisation length is halved?",
         answer="Keep the mesh-to-length-scale ratio comparable: halving the length scale suggests halving the local mesh spacing, followed by a resolution study.",
         scope="Analytic isolated AT2 profile. Lengths use one arbitrary unit.",
         source="source/animations/mechanics_scenes.py: PhaseFieldBand"),
    dict(key="explicit_implicit_step", title="Follow one partitioned dynamic update",
         movie="assets/animations/mechanics/explicit_implicit_step.mp4",
         poster="assets/animations/mechanics/explicit_implicit_step_poster.png",
         section="L1 · 26–40 min", pause=11.0,
         point="Mechanics\n\nDriving energy\n\nDamage\n\nAccept the checked state",
         question="Which fields are held fixed during the mechanics and damage updates?",
         answer="The explicit mechanics update uses the accepted damage. Its updated tensile energy and stored history drive the implicit damage problem. Describe the time-step stability and acceptance checks before advancing.",
         scope="Partitioned dynamic schematic. The PhAST practical uses quasistatic mechanics.",
         source="source/animations/mechanics_scenes.py: ExplicitImplicitStep"),
    dict(key="reverse_accumulation", title="Every use of a parameter contributes",
         movie="assets/animations/learning/reverse_accumulation.mp4",
         poster="assets/animations/learning/reverse_accumulation.png",
         section="L2 · 22–35 min", pause=15.0,
         point="Clip notation\n\nx: state z\nθ: parameter p\nL: loss J\n\nSum both parameter contributions",
         question="Why does the final derivative contain one contribution from each update?",
         answer="The shared parameter enters both update maps. Reverse accumulation propagates the terminal loss sensitivity through the state path and adds the local parameter-Jacobian products at both uses. The clip uses x, theta and L for the state, parameter and loss denoted z, p and J in the preceding slides.",
         scope="Two differentiable updates; fixed initial state; terminal scalar loss.",
         source="source/animations/learning_scenes.py: ReverseAccumulation"),
    dict(key="history_switch", title="Stored history selects the active sensitivity",
         movie="source/book/research/visuals/history_switch.mp4",
         poster="source/book/research/visuals/history_switch.png",
         section="L2 · 35–45 min", pause=10.0,
         point="Energy can unload.\nHistory retains its peak.\n\nThe active branch determines the local derivative.\n\nAt a tie, specify the derivative convention.",
         question="During unloading below the previous peak, which input receives the hard-history sensitivity?",
         answer="For H_new=max(H_old, psi), the old history receives the derivative when psi<H_old; the current energy receives it when psi>H_old. At equality the maximum is nonsmooth. The displayed 0.5 marker is a selected tie convention. The green sigmoid weights define a surrogate backward rule; they describe a different derivative rule from the hard maximum away from the branch transition as well.",
         scope="Dimensionless material-point teaching example. Green: surrogate reverse weights.",
         source="source/book/research/code/history_animation.py; history_lesson.py"),
    dict(key="checked_learned_proposal", title="A learned field enters a checked interface",
         movie="assets/animations/learning/checked_learned_proposal.mp4",
         poster="assets/animations/learning/checked_learned_proposal.png",
         section="L3 · 32–45 min", pause=15.0,
         point="Declare features.\n\nPropose a field.\n\nCheck and correct.\n\nRecheck the result.",
         question="Which checks should a proposed damage field satisfy before the simulation accepts it?",
         answer="Check the declared node/graph ordering, units, normalisation and output location first. Check prescribed values, bounds, irreversibility and the declared residual tolerance. A corrected state is checked again. The graphic names Radius GNO/GNN as possible architecture contracts; the practical trains MLP/RBF models for a Helmholtz-type field problem. Total cost includes graph construction, prediction, conversion, checks and correction.",
         scope="Interface schematic. Practical models: MLP/RBF on a Helmholtz-type teaching field.",
         source="source/animations/learning_scenes.py: CheckedLearnedProposal"),
]


def sha(p):
    return hashlib.sha256(Path(p).read_bytes()).hexdigest()


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    media = OUT / "media"
    media.mkdir(exist_ok=True)
    p = Presentation(SOURCE)
    assert len(p.slides) == 18
    original_hash = sha(SOURCE)
    scale = p.slide_width / 2560

    def pos(value):
        return round(value * scale)

    def textbox(sl, text, x, y, w, h, size=48, color=INK, bold=False):
        shape = sl.shapes.add_textbox(pos(x), pos(y), pos(w), pos(h))
        shape.text_frame.word_wrap = True
        shape.text_frame.margin_left = shape.text_frame.margin_right = 0
        shape.text_frame.margin_top = shape.text_frame.margin_bottom = 0
        shape.text = text
        for paragraph in shape.text_frame.paragraphs:
            paragraph.font.name = "Arial"
            paragraph.font.size = Pt(size)
            paragraph.font.bold = bold
            paragraph.font.color.rgb = RGBColor.from_string(color)
        return shape

    # Keep inherited body text, equations and connectors editable. Apply the
    # requested white/orange/blue design and a reproducible available font.
    for sl in p.slides:
        sl.background.fill.solid()
        sl.background.fill.fore_color.rgb = RGBColor(255, 255, 255)
        for shape in sl.shapes:
            if not shape.has_text_frame:
                continue
            replacements = {"Forward calculation": "Forward solve",
                            "Backward calculation": "Adjoint solve"}
            if shape.text in replacements:
                shape.text_frame.paragraphs[0].runs[0].text = replacements[shape.text]
            for para in shape.text_frame.paragraphs:
                para.font.name = "Arial"
                for run in para.runs:
                    run.font.name = "Arial"
                    if run.font.color.type is not None and run.font.color.type == 1:
                        if str(run.font.color.rgb).upper() == "B91862":
                            run.font.color.rgb = RGBColor.from_string(BLUE)
    notes_by_original = {}
    for i, sl in enumerate(p.slides, 1):
        notes_by_original[i] = sl.notes_slide.notes_text_frame.text

    clip_ids = {}
    receipts = []
    for clip in CLIPS:
        src = ROOT / clip["movie"]
        dst = media / src.name
        poster = media / f"{clip['key']}_poster.png"
        shutil.copyfile(ROOT / clip["poster"], poster)
        probe = json.loads(subprocess.check_output([
            "ffprobe", "-v", "error", "-select_streams", "v:0",
            "-show_entries", "stream=codec_name,profile,pix_fmt,width,height,r_frame_rate",
            "-show_entries", "format=duration", "-of", "json", str(src)]))
        # The history primitive is 0.5 fps. A 30-fps H.264 copy retains each
        # original frame for two seconds and improves native-player seeking.
        if clip["key"] == "history_switch":
            subprocess.run(["ffmpeg", "-y", "-v", "error", "-i", str(src),
                            "-vf", "fps=30", "-c:v", "libx264", "-crf", "18",
                            "-pix_fmt", "yuv420p", "-movflags", "+faststart", str(dst)], check=True)
        else:
            shutil.copyfile(src, dst)
        subprocess.run(["ffmpeg", "-v", "error", "-i", str(dst), "-f", "null", "-"], check=True)
        sl = p.slides.add_slide(p.slide_layouts[16])
        for shape in list(sl.shapes):
            if shape.is_placeholder:
                element = shape._element
                element.getparent().remove(element)
        sl.background.fill.solid()
        sl.background.fill.fore_color.rgb = RGBColor(255, 255, 255)
        textbox(sl, clip["title"], 128, 90, 2304, 145, 70, BLUE, True)
        width, height = probe["streams"][0]["width"], probe["streams"][0]["height"]
        mh = 990
        mw = mh * width / height
        left, top = 100, 285
        # A native picture beneath the movie is a static fallback for PDF
        # exporters and viewers that hide the movie object entirely.
        still = sl.shapes.add_picture(str(poster), pos(left), pos(top), pos(mw), pos(mh))
        still.name = f"Static fallback: {clip['key']}"
        movie = sl.shapes.add_movie(str(dst), pos(left), pos(top), pos(mw), pos(mh),
                                   poster_frame_image=str(poster), mime_type="video/mp4")
        movie.name = f"Play: {clip['key']}"
        movie._element.xpath(".//p:cNvPr")[0].set("descr", clip["scope"] + " " + clip["question"])
        tx = left + mw + 85
        textbox(sl, "Observe", tx, 335, 2400-tx, 80, 46, ORANGE, True)
        textbox(sl, clip["point"], tx, 455, 2400-tx, 695, 45, INK)
        textbox(sl, "Click the figure to play or pause.", 128, 1288, 2240, 48, 29, MUTED)
        textbox(sl, clip["scope"], 128, 1352, 2170, 46, 26, MUTED)
        duration = float(probe["format"]["duration"])
        sl.notes_slide.notes_text_frame.text = (
            f"{clip['section']}. {clip['scope']}\n\n"
            f"Play the {duration:g}-second clip. Pause around {clip['pause']:g} seconds.\n"
            f"Ask: {clip['question']}\n\nDiscussion: {clip['answer']}\n\n"
            f"Static fallback: media/{poster.name}. Describe the same sequence from its poster.\n"
            "Replay from the beginning after discussion. The clip is embedded in this deck and needs no network connection.\n\n"
            f"Original source: {clip['source']}\n\n{CREDIT}")
        clip_ids[clip["key"]] = len(p.slides)
        receipts.append(dict(key=clip["key"], source=clip["movie"], source_sha256=sha(src),
                             embedded_copy=f"media/{dst.name}", embedded_sha256=sha(dst),
                             poster=f"media/{poster.name}", poster_sha256=sha(poster),
                             duration_seconds=duration, original_probe=probe,
                             pause_seconds=clip["pause"], teaching_scope=clip["scope"]))

    # One coherent introductory spine: L1 mechanics, L2 derivatives, L3 learning.
    # All 18 retained slides appear exactly once; five movie slides are inserted.
    order = [1, 5, 6, clip_ids["phase_field_band"], 7, 8, clip_ids["explicit_implicit_step"],
             9, 15, 16, 10, 11, clip_ids["reverse_accumulation"], 12,
             clip_ids["history_switch"], 13, 14, 2, 3, 4,
             clip_ids["checked_learned_proposal"], 17, 18]
    assert sorted(order) == list(range(1, 24))
    ids = list(p.slides._sldIdLst)
    for node in ids:
        p.slides._sldIdLst.remove(node)
    for i in order:
        p.slides._sldIdLst.append(ids[i-1])

    mapping = []
    for number, (old, sl) in enumerate(zip(order, p.slides), 1):
        section = "L1" if number <= 10 else "L2" if number <= 17 else "L3 / practical bridge"
        for shape in sl.shapes:
            if shape.has_text_frame and re.fullmatch(r"\d{2}", shape.text.strip()):
                shape.text_frame.paragraphs[0].runs[0].text = f"{number:02d}"
        if old >= 19:
            textbox(sl, f"{number:02d}", 2330, 1352, 105, 46, 28, MUTED)
        if old <= 18:
            sl.notes_slide.notes_text_frame.text = (
                f"Delivery placement: {section}.\n\n" + notes_by_original[old])
        title = next((s.text for s in sl.shapes if s.has_text_frame and
                      s.top < pos(400) and len(s.text.strip()) > 5), "")
        mapping.append(dict(slide=number, section=section, title=title,
                            original_v7_slide=old if old <= 18 else None,
                            clip=next((k for k, v in clip_ids.items() if v == old), None)))

    p.core_properties.title = "PhAST: Numerical methods and deep learning — animated lecture spine"
    p.core_properties.author = "Allamaprabhu Ani"
    p.core_properties.subject = ("UKACM Autumn School 2026; presented by Sathiskumar A. Ponnusami, "
                                 "Queen Mary University of London; CEMS-Lab")
    p.core_properties.comments = CREDIT
    p.core_properties.last_modified_by = "CEMS-Lab"
    target = OUT / "phast_lecture_animated.pptx"
    p.save(target)
    assert sha(SOURCE) == original_hash
    receipt = dict(source_deck=str(SOURCE.relative_to(ROOT)), source_sha256=original_hash,
                   candidate=str(target.relative_to(ROOT)), candidate_sha256=sha(target),
                   slides=23, media_count=5, mapping=mapping, media=receipts,
                   native_playback="Requires presentation-machine testing in PowerPoint and Keynote.")
    (OUT / "media_manifest.json").write_text(json.dumps(receipt, indent=2) + "\n")
    (OUT / "speaker_notes.md").write_text("# Animated lecture speaker notes\n\n" + "\n".join(
        f"## {item['slide']}. {item['title']}\n\n{sl.notes_slide.notes_text_frame.text}\n"
        for item, sl in zip(mapping, p.slides)))
    print(json.dumps(dict(output=str(target), slides=23, embedded_clips=5), indent=2))


if __name__ == "__main__":
    main()
