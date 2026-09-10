"""Check the animated candidate's portability, provenance and editability."""
import hashlib
import json
import zipfile
from pathlib import Path

from lxml import etree
from pptx import Presentation

ROOT = Path(__file__).resolve().parents[3]
OUT = ROOT / "output/lecture-20260910"
manifest = json.loads((OUT / "media_manifest.json").read_text())
deck = ROOT / manifest["candidate"]
sha = lambda blob: hashlib.sha256(blob).hexdigest()
p = Presentation(deck)
with zipfile.ZipFile(deck) as package:
    embedded = {name: sha(package.read(name)) for name in package.namelist()
                if name.startswith("ppt/media/") and name.endswith(".mp4")}
    external_movies = []
    for name in package.namelist():
        if name.endswith(".rels"):
            for rel in etree.fromstring(package.read(name)):
                if rel.get("TargetMode") == "External" and any(
                        kind in rel.get("Type", "") for kind in ("video", "media")):
                    external_movies.append((name, rel.get("Target")))
out_of_bounds = [(i, shape.name) for i, slide in enumerate(p.slides, 1)
                 for shape in slide.shapes if shape.left < 0 or shape.top < 0
                 or shape.left + shape.width > p.slide_width + 10
                 or shape.top + shape.height > p.slide_height + 10]
checks = {
    "23_slides": len(p.slides) == 23,
    "five_internal_movies": len(embedded) == 5 and not external_movies,
    "all_movies_match_manifest": set(embedded.values()) == {
        clip["embedded_sha256"] for clip in manifest["media"]},
    "five_native_static_posters": sum(shape.name.startswith("Static fallback:")
        for slide in p.slides for shape in slide.shapes) == 5,
    "all_slides_have_presenter_notes": all(slide.notes_slide.notes_text_frame.text
                                           for slide in p.slides),
    "movie_slides_have_pause_questions": all(
        "Pause around" in p.slides[item["slide"] - 1].notes_slide.notes_text_frame.text
        for item in manifest["mapping"] if item["clip"]),
    "no_shapes_outside_slide": not out_of_bounds,
    "source_v7_unchanged": sha((ROOT / manifest["source_deck"]).read_bytes()) ==
        manifest["source_sha256"],
    "candidate_matches_manifest": sha(deck.read_bytes()) == manifest["candidate_sha256"],
    "creator_metadata": p.core_properties.author == "Allamaprabhu Ani",
}
result = {"checks": checks, "all_passed": all(checks.values()),
          "candidate_sha256": sha(deck.read_bytes()), "embedded_movies": embedded,
          "external_movies": external_movies, "out_of_bounds": out_of_bounds,
          "native_playback": "Separate PowerPoint/Keynote presentation-machine check required."}
(OUT / "package_checks.json").write_text(json.dumps(result, indent=2) + "\n")
print(json.dumps(result, indent=2))
assert result["all_passed"]
